"""SQLite terminology store with exact and BM25 description retrieval."""

from __future__ import annotations

import os
import sqlite3
import hashlib
import csv
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .data import normalize

FILES = {
    "ICD10CM": "icd10cm_codes.xlsx",
    "SNOMEDCT_US": "SNOMEDCT_US_codes_and_descriptions.xlsx",
}
TABLE = {"ICD10CM": "icd", "SNOMEDCT_US": "snomed"}


def _source_signature(path: Path) -> str:
    stat = path.stat()
    return f"{path.name}:{stat.st_size}:{stat.st_mtime_ns}"


def _index_code_signature() -> str:
    digest = hashlib.sha256()
    for filename in ("data.py", "terminology.py"):
        digest.update((Path(__file__).parent / filename).read_bytes())
    return digest.hexdigest()


class Terminology:
    def __init__(self, path: str | Path):
        self.path = Path(path)
        self.conn = sqlite3.connect(f"file:{self.path}?mode=ro", uri=True)
        self.conn.row_factory = sqlite3.Row

    def close(self) -> None:
        self.conn.close()

    def __enter__(self) -> "Terminology":
        return self

    def __exit__(self, *_: object) -> None:
        self.close()

    def source_signatures(self) -> dict[str, str]:
        return dict(self.conn.execute("SELECT key, value FROM metadata"))

    def contains(self, system: str, code: str) -> bool:
        table = TABLE[system]
        return self.conn.execute(f"SELECT 1 FROM {table}_codes WHERE code=?", (code,)).fetchone() is not None

    def descriptions(self, system: str, codes: Iterable[str], limit_per_code: int = 30) -> dict[str, list[str]]:
        codes = list(dict.fromkeys(codes))
        result: dict[str, list[str]] = {code: [] for code in codes}
        if not codes:
            return result
        table = TABLE[system]
        for start in range(0, len(codes), 400):
            part = codes[start:start + 400]
            placeholders = ",".join("?" for _ in part)
            query = f"SELECT code,description FROM {table}_docs WHERE code IN ({placeholders}) ORDER BY id"
            for row in self.conn.execute(query, part):
                if len(result[row["code"]]) < limit_per_code:
                    result[row["code"]].append(row["description"])
        return result

    def iter_codes(self, system: str):
        """Yield legal code IDs and their first supplied description in stable order."""
        table = TABLE[system]
        for row in self.conn.execute(f"SELECT code,description FROM {table}_codes ORDER BY code"):
            yield row["code"], row["description"]

    def iter_dense_descriptions(self, system: str, max_per_code: int = 2,
                                limit_codes: int | None = None):
        """Yield the first distinct normalized descriptions for each legal code."""
        if max_per_code < 1 or (limit_codes is not None and limit_codes < 1):
            raise ValueError("Description and code limits must be positive")
        table = TABLE[system]
        query = (f"SELECT d.code,d.description,d.norm FROM {table}_docs d "
                 f"JOIN {table}_codes c ON c.code=d.code ORDER BY d.code,d.id")
        current_code = None
        code_count = description_count = 0
        seen_norms: set[str] = set()
        for row in self.conn.execute(query):
            code = row["code"]
            if code != current_code:
                code_count += 1
                if limit_codes is not None and code_count > limit_codes:
                    break
                current_code = code
                description_count = 0
                seen_norms.clear()
            norm = row["norm"]
            if description_count >= max_per_code or norm in seen_norms:
                continue
            seen_norms.add(norm)
            description_count += 1
            yield code, row["description"]

    def snomed_neighbors(self, codes: Iterable[str], limit_per_code: int = 20) -> dict[str, list[str]]:
        """Return legal RF2 is-a parents and children if a licensed snapshot was indexed."""
        codes = list(dict.fromkeys(codes))
        result = {code: [] for code in codes}
        if not codes:
            return result
        if self.conn.execute("SELECT 1 FROM sqlite_master WHERE name='snomed_isa'").fetchone() is None:
            return result
        for code in codes:
            rows = self.conn.execute(
                "SELECT parent AS neighbor FROM snomed_isa WHERE child=? "
                "UNION SELECT child AS neighbor FROM snomed_isa WHERE parent=? "
                "ORDER BY neighbor LIMIT ?", (code, code, limit_per_code))
            result[code] = [row["neighbor"] for row in rows]
        return result

    def umls_links(self, system: str, codes: Iterable[str], limit_per_code: int = 20) -> dict[str, list[str]]:
        """Return CUI-linked codes in the other system as a weak prior, never a crosswalk truth."""
        codes = list(dict.fromkeys(codes))
        result = {code: [] for code in codes}
        if self.conn.execute("SELECT 1 FROM sqlite_master WHERE name='umls_pairs'").fetchone() is None:
            return result
        source, target = ("icd", "snomed") if system == "ICD10CM" else ("snomed", "icd")
        for code in codes:
            rows = self.conn.execute(
                f"SELECT {target} FROM umls_pairs WHERE {source}=? ORDER BY {target} LIMIT ?",
                (code, limit_per_code))
            result[code] = [row[0] for row in rows]
        return result

    def search(self, system: str, term: str, limit: int = 100) -> dict[str, tuple[float, str]]:
        """Return code -> (retrieval strength, matched description)."""
        table = TABLE[system]
        norm = normalize(term)
        if not norm:
            return {}
        result: dict[str, tuple[float, str]] = {}
        for row in self.conn.execute(f"SELECT code,description FROM {table}_docs WHERE norm=? ORDER BY id LIMIT 300", (norm,)):
            result[row["code"]] = (1.0, row["description"])

        tokens = list(dict.fromkeys(t for t in norm.split() if len(t) > 1 and not t.isdigit()))
        if not tokens:
            return result
        tokens = tokens[:10]
        quoted = [f'"{t.replace(chr(34), "")}"' for t in tokens]
        queries = [(f'"{" ".join(tokens)}"', 0.95)]
        if len(tokens) > 1:
            queries.append((" AND ".join(quoted), 0.78))
            queries.append((" OR ".join(quoted), 0.48))
        else:
            queries.append((quoted[0] + "*", 0.55))
        for expression, weight in queries:
            sql = (f"SELECT d.code,d.description FROM {table}_fts f "
                   f"JOIN {table}_docs d ON d.id=f.rowid "
                   f"WHERE {table}_fts MATCH ? ORDER BY bm25({table}_fts),d.code,d.id LIMIT ?")
            try:
                for rank, row in enumerate(self.conn.execute(sql, (expression, limit * 2))):
                    strength = weight / (1.0 + rank / 20.0)
                    old = result.get(row["code"])
                    if old is None or strength > old[0]:
                        result[row["code"]] = (strength, row["description"])
            except sqlite3.OperationalError:
                # A malformed source token cannot prevent the other retrieval paths.
                continue
        return dict(sorted(result.items(), key=lambda kv: (-kv[1][0], kv[0]))[:limit])


def build_terminology(data_dir: str | Path, output: str | Path, force: bool = False,
                      umls_mrconso: str | Path | None = None,
                      snomed_rf2_descriptions: str | Path | None = None,
                      snomed_rf2_relationships: str | Path | None = None) -> Path:
    data_dir, output = Path(data_dir), Path(output)
    signatures = {name: _source_signature(data_dir / filename) for name, filename in FILES.items()}
    licensed = {"umls_mrconso": umls_mrconso,
                "snomed_rf2_descriptions": snomed_rf2_descriptions,
                "snomed_rf2_relationships": snomed_rf2_relationships}
    signatures.update({name: _source_signature(Path(path)) for name, path in licensed.items() if path})
    signatures["index_code_sha256"] = _index_code_signature()
    if output.exists() and not force:
        try:
            with Terminology(output) as existing:
                if existing.source_signatures() == signatures:
                    return output
        except sqlite3.DatabaseError:
            pass

    output.parent.mkdir(parents=True, exist_ok=True)
    temp = output.with_name(output.name + ".building")
    temp.unlink(missing_ok=True)
    conn = sqlite3.connect(temp)
    try:
        conn.execute("PRAGMA journal_mode=OFF")
        conn.execute("PRAGMA synchronous=OFF")
        conn.execute("PRAGMA temp_store=MEMORY")
        conn.execute("PRAGMA cache_size=-200000")
        conn.execute("CREATE TABLE metadata(key TEXT PRIMARY KEY,value TEXT NOT NULL)")
        conn.executemany("INSERT INTO metadata VALUES (?,?)", signatures.items())
        for table in TABLE.values():
            conn.execute(f"CREATE TABLE {table}_docs(id INTEGER PRIMARY KEY,code TEXT NOT NULL,description TEXT NOT NULL,norm TEXT NOT NULL)")
            conn.execute(f"CREATE TABLE {table}_codes(code TEXT PRIMARY KEY,description TEXT NOT NULL)")
            conn.execute(f"CREATE VIRTUAL TABLE {table}_fts USING fts5(description,content='{table}_docs',content_rowid='id',tokenize='unicode61 remove_diacritics 2')")

        legal_codes: dict[str, set[str]] = {}

        for system, filename in FILES.items():
            table = TABLE[system]
            workbook = load_workbook(data_dir / filename, read_only=True, data_only=True)
            sheet = workbook.worksheets[0]
            batch: list[tuple[str, str, str]] = []
            count = 0
            for row in sheet.iter_rows(min_row=2, values_only=True):
                code = str(row[0]).strip() if row[0] is not None else ""
                description = str(row[1]).strip() if row[1] is not None else ""
                if not code or not description:
                    continue
                batch.append((code, description, normalize(description)))
                if len(batch) >= 10_000:
                    _insert_batch(conn, table, batch)
                    count += len(batch)
                    batch.clear()
                    if count % 100_000 == 0:
                        print(f"Indexed {system}: {count:,} descriptions", flush=True)
            if batch:
                _insert_batch(conn, table, batch)
                count += len(batch)
            workbook.close()
            legal_codes[system] = {row[0] for row in conn.execute(f"SELECT code FROM {table}_codes")}
            print(f"Indexed {system}: {count:,} supplied descriptions", flush=True)

        if umls_mrconso:
            _import_umls(conn, Path(umls_mrconso), legal_codes)
        if snomed_rf2_descriptions:
            _import_snomed_descriptions(conn, Path(snomed_rf2_descriptions), legal_codes["SNOMEDCT_US"])
        if snomed_rf2_relationships:
            _import_snomed_relationships(conn, Path(snomed_rf2_relationships), legal_codes["SNOMEDCT_US"])
        for table in TABLE.values():
            conn.execute(f"CREATE INDEX {table}_norm_idx ON {table}_docs(norm)")
            conn.execute(f"CREATE INDEX {table}_code_idx ON {table}_docs(code)")
            conn.execute(f"INSERT INTO {table}_fts({table}_fts) VALUES ('rebuild')")
            conn.commit()
        conn.execute("PRAGMA optimize")
        conn.commit()
    except BaseException:
        conn.close()
        temp.unlink(missing_ok=True)
        raise
    conn.close()
    os.replace(temp, output)
    return output


def _insert_batch(conn: sqlite3.Connection, table: str, rows: list[tuple[str, str, str]]) -> None:
    conn.executemany(f"INSERT INTO {table}_docs(code,description,norm) VALUES (?,?,?)", rows)
    conn.executemany(f"INSERT OR IGNORE INTO {table}_codes(code,description) VALUES (?,?)",
                     ((code, desc) for code, desc, _ in rows))
    conn.commit()


def _import_umls(conn: sqlite3.Connection, path: Path, legal: dict[str, set[str]]) -> None:
    """Import English MRCONSO atoms only for codes in the supplied answer space."""
    batches: dict[str, list[tuple[str, str, str]]] = {system: [] for system in TABLE}
    seen: set[tuple[str, str, str]] = set()
    cui_codes: dict[str, dict[str, set[str]]] = {}
    counts = {system: 0 for system in TABLE}
    with path.open(encoding="utf-8", errors="replace") as handle:
        for line in handle:
            row = line.rstrip("\n").split("|")
            if len(row) < 17 or row[1] != "ENG" or row[16] != "N":
                continue
            system = row[11]
            if system not in TABLE:
                continue
            code, description = row[13].strip(), row[14].strip()
            if system == "ICD10CM" and code not in legal[system] and "." not in code and len(code) > 3:
                code = code[:3] + "." + code[3:]
            if code not in legal[system] or not description:
                continue
            cui_codes.setdefault(row[0], {}).setdefault(system, set()).add(code)
            norm = normalize(description)
            key = (system, code, norm)
            if not norm or key in seen:
                continue
            seen.add(key)
            batches[system].append((code, description, norm))
            if len(batches[system]) >= 10_000:
                _insert_batch(conn, TABLE[system], batches[system])
                counts[system] += len(batches[system])
                batches[system].clear()
    for system, batch in batches.items():
        if batch:
            _insert_batch(conn, TABLE[system], batch)
            counts[system] += len(batch)
    conn.execute("CREATE TABLE umls_pairs(icd TEXT NOT NULL,snomed TEXT NOT NULL,PRIMARY KEY(icd,snomed))")
    pair_batch: list[tuple[str, str]] = []
    for linked in cui_codes.values():
        icd = linked.get("ICD10CM", set())
        snomed = linked.get("SNOMEDCT_US", set())
        if len(icd) > 10 or len(snomed) > 10:
            continue
        for icd_code in sorted(icd):
            for snomed_code in sorted(snomed):
                pair_batch.append((icd_code, snomed_code))
                if len(pair_batch) >= 10_000:
                    conn.executemany("INSERT OR IGNORE INTO umls_pairs VALUES (?,?)", pair_batch)
                    pair_batch.clear()
    if pair_batch:
        conn.executemany("INSERT OR IGNORE INTO umls_pairs VALUES (?,?)", pair_batch)
    conn.execute("CREATE INDEX umls_pairs_snomed_idx ON umls_pairs(snomed)")
    conn.commit()
    print(f"Imported UMLS aliases: {counts}", flush=True)


def _import_snomed_descriptions(conn: sqlite3.Connection, path: Path, legal: set[str]) -> None:
    batch: list[tuple[str, str, str]] = []
    count = 0
    with path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        required = {"active", "conceptId", "term"}
        if not required.issubset(reader.fieldnames or []):
            raise ValueError(f"SNOMED RF2 description snapshot needs {sorted(required)}")
        for row in reader:
            code, description = row["conceptId"], row["term"].strip()
            if row["active"] != "1" or code not in legal or not description:
                continue
            batch.append((code, description, normalize(description)))
            if len(batch) >= 10_000:
                _insert_batch(conn, "snomed", batch)
                count += len(batch)
                batch.clear()
    if batch:
        _insert_batch(conn, "snomed", batch)
        count += len(batch)
    print(f"Imported {count:,} active SNOMED RF2 descriptions", flush=True)


def _import_snomed_relationships(conn: sqlite3.Connection, path: Path, legal: set[str]) -> None:
    conn.execute("CREATE TABLE snomed_isa(child TEXT NOT NULL,parent TEXT NOT NULL,PRIMARY KEY(child,parent))")
    batch: list[tuple[str, str]] = []
    count = 0
    with path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        required = {"active", "sourceId", "destinationId", "typeId"}
        if not required.issubset(reader.fieldnames or []):
            raise ValueError(f"SNOMED RF2 relationship snapshot needs {sorted(required)}")
        for row in reader:
            child, parent = row["sourceId"], row["destinationId"]
            if row["active"] != "1" or row["typeId"] != "116680003":
                continue
            if child not in legal or parent not in legal:
                continue
            batch.append((child, parent))
            if len(batch) >= 10_000:
                conn.executemany("INSERT OR IGNORE INTO snomed_isa VALUES (?,?)", batch)
                count += len(batch)
                batch.clear()
    if batch:
        conn.executemany("INSERT OR IGNORE INTO snomed_isa VALUES (?,?)", batch)
        count += len(batch)
    conn.execute("CREATE INDEX snomed_isa_parent_idx ON snomed_isa(parent)")
    conn.commit()
    print(f"Imported {count:,} active SNOMED is-a relationships", flush=True)
